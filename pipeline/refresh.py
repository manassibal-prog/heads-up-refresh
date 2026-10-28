"""Heads-up call list refresher.

Runs Adit's query (unchanged except start_date) through the Metabase API and
keeps two tabs up to date WITHOUT ever moving a row:
  - caller sheet  : CALL-arm tickets only (what agents see)
  - DS register   : both arms (permanent record for the analysis)

Rules for the sheets:
  - new tickets are appended at the bottom, oldest flag first
  - existing rows are matched by ticket_id and only their live columns
    (status etc.) are overwritten; agent-typed columns are never touched
  - nothing is sorted or deleted
"""
import json
import os
import pathlib
import re
import sys
import urllib.request
from datetime import datetime, timedelta, timezone

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
IST = timezone(timedelta(hours=5, minutes=30))
MODEL = "9f70571979"  # model prefix hard-coded in the SQL

# caller sheet: A-F are filled once on append, G-P belong to the agents,
# Q-S are live columns refreshed every run
CALLER_FIXED = ["TICKET_ID", "KAPTURE_TICKET_ID", "CSP_ID", "CSP_NAME", "FLAGGED_AT_IST", "SCORE_BAND"]
CALLER_LIVE_HDR = ["status_now", "issue", "last_refreshed"]
CALLER_LIVE_COL = 17  # Q
CALLER_ADDED_COL = 20  # T: added_at, written once when the ticket is first added to the sheet

REG_HDR = ["ticket_id", "kapture_ticket_id", "csp_id", "csp_name", "arm", "score_band", "delay_score",
           "model", "created_at_ist", "flagged_at_ist", "enrolled_at", "status_now", "status_updated_at", "issue"]
REG_LIVE_COL = 12  # L: status_now, status_updated_at


def log(msg):
    line = f"{datetime.now(IST):%Y-%m-%d %H:%M:%S} {msg}"
    if sys.stdout:  # pythonw (no console) has no stdout
        print(line)
    with open(HERE / "refresh.log", "a", encoding="utf-8") as f:
        f.write(line + "\n")


def load_cfg():
    cfg = json.load(open(HERE / "config.json", encoding="utf-8"))
    if os.environ.get("TRIAL"):  # trial: today's data into temp tabs of the register sheet
        rid = cfg["register"]["sheet_id"]
        cfg["start_date"] = os.environ["TRIAL"]
        cfg["caller"] = {"sheet_id": rid, "tab": "tmp_caller"}
        cfg["register"] = {"sheet_id": rid, "tab": "tmp_reg"}
    if os.environ.get("MB_KEY"):  # cloud: keys come from secrets as env vars
        cfg["mb_url"], cfg["mb_key"] = os.environ["MB_URL"].rstrip("/"), os.environ["MB_KEY"]
    else:
        env = (ROOT / "metabase.env").read_text(encoding="utf-8").splitlines()
        cfg["mb_url"] = env[0].strip().rstrip("/")
        cfg["mb_key"] = next(l.split(":", 1)[1].strip() for l in reversed(env) if l.startswith(("key:", "key_secret:")))
    return cfg


def fmt_ts(v):
    if not v:
        return ""
    s = str(v).replace("T", " ")
    return re.sub(r"(\.\d+)?(Z|[+-]\d\d:\d\d)?$", "", s)[:16]


def fetch(cfg):
    sql = (HERE / "heads_up_call_list.sql").read_text(encoding="utf-8")
    sql, n = re.subn(r"'\d{4}-\d{2}-\d{2}'::date AS start_date", f"'{cfg['start_date']}'::date AS start_date", sql)
    assert n == 1, "start_date line not found in SQL"
    body = json.dumps({"database": 113, "type": "native", "native": {"query": sql}}).encode()
    req = urllib.request.Request(cfg["mb_url"] + "/api/dataset", body,
                                 {"x-api-key": cfg["mb_key"], "Content-Type": "application/json"})
    for attempt in range(3):
        try:
            d = json.load(urllib.request.urlopen(req, timeout=120))
            break
        except (TimeoutError, OSError) as e:
            log(f"Metabase attempt {attempt + 1} failed: {e}")
            if attempt == 2:
                raise
    if d.get("error") or d.get("status") == "failed":
        raise RuntimeError(f"Metabase error: {str(d.get('error'))[:300]}")
    cols = [c["name"] for c in d["data"]["cols"]]
    rows = [dict(zip(cols, r)) for r in d["data"]["rows"]]
    for r in rows:
        r["TICKET_ID"] = str(r["TICKET_ID"])
        r["CREATED_AT_IST"] = fmt_ts(r["CREATED_AT_IST"])
        r["FLAGGED_AT_IST"] = fmt_ts(r["FLAGGED_AT_IST"])
        r["KAPTURE_TICKET_ID"] = "" if r["KAPTURE_TICKET_ID"] is None else str(r["KAPTURE_TICKET_ID"])
        for k in ("CSP_NAME", "ISSUE"):
            r[k] = r[k] or ""
    fill_missing_kapture(cfg, rows)
    rows.sort(key=lambda r: (r["FLAGGED_AT_IST"], r["TICKET_ID"]))
    return rows


def fill_missing_kapture(cfg, rows):
    """service_ticket_model lags upstream: where its kapture_ticket_id is blank, look it up in
    the raw tickets table. Adit's query is not touched; failures here are non-fatal."""
    miss = [r for r in rows if not r["KAPTURE_TICKET_ID"]]
    if not miss:
        return
    ids = ",".join(f"'{r['TICKET_ID']}'" for r in miss if r["TICKET_ID"].isdigit())
    sql = ("select ticket_id::text, max(kapture_ticket_id)::text from prod_db.dynamodb_read.tickets "
           f"where ticket_id::text in ({ids}) and kapture_ticket_id is not null group by 1")
    try:
        body = json.dumps({"database": 113, "type": "native", "native": {"query": sql}}).encode()
        req = urllib.request.Request(cfg["mb_url"] + "/api/dataset", body,
                                     {"x-api-key": cfg["mb_key"], "Content-Type": "application/json"})
        d = json.load(urllib.request.urlopen(req, timeout=120))
        found = {a: b for a, b in d["data"]["rows"]}
    except Exception as e:
        log(f"kapture lookup failed (non-fatal): {e}")
        return
    for r in miss:
        r["KAPTURE_TICKET_ID"] = found.get(r["TICKET_ID"], "")
    log(f"kapture id missing for {len(miss)} tickets, filled {sum(1 for r in miss if r['KAPTURE_TICKET_ID'])} from upstream")


# ---------- storage backends: read_all() and update_block() only ----------
class GTab:
    def __init__(self, sheet_id, tab, sa_file):
        import gspread
        if os.environ.get("GOOGLE_SA_JSON"):
            gc = gspread.service_account_from_dict(json.loads(os.environ["GOOGLE_SA_JSON"]))
        else:
            gc = gspread.service_account(filename=str(ROOT / sa_file))
        self.ws = gc.open_by_key(sheet_id).worksheet(tab)

    def read_all(self):
        return self.ws.get_all_values()

    def update_block(self, row, col, values):
        if not values:
            return
        from gspread.utils import rowcol_to_a1
        rng = f"{rowcol_to_a1(row, col)}:{rowcol_to_a1(row + len(values) - 1, col + len(values[0]) - 1)}"
        self.ws.update(range_name=rng, values=values, value_input_option="RAW")


class XTab:
    """Local .xlsx fallback / test backend."""

    def __init__(self, path, tab):
        import openpyxl
        self.path, self.tab = ROOT / path, tab
        self.wb = openpyxl.load_workbook(self.path)
        self.ws = self.wb[tab]

    def read_all(self):
        return [["" if c is None else str(c) for c in r] for r in self.ws.iter_rows(values_only=True)]

    def update_block(self, row, col, values):
        for i, vals in enumerate(values):
            for j, v in enumerate(vals):
                self.ws.cell(row=row + i, column=col + j, value=v)
        self.wb.save(self.path)


def open_tab(cfg, which):
    c = cfg[which]
    if cfg["backend"] == "gsheets":
        return GTab(c["sheet_id"], c["tab"], cfg["service_account_file"])
    return XTab(c["xlsx"], c["tab"])


def nid(v):
    """ticket ids pasted by hand may carry commas/spaces ('1,792,727,978,487,000'): compare digits only"""
    return re.sub(r"\D", "", str(v))


def next_row(rows):
    n = 1
    for i, r in enumerate(rows, 1):
        if r and r[0].strip():
            n = i
    return n + 1


def sync(tab, fresh, fixed_vals, live_vals, live_col, live_hdr, extra=None):
    """fresh: rows to hold; fixed_vals(r)->list written once at append;
    live_vals(r)->list refreshed every run in live_col onward.
    extra=(col, header, value): a write-once column stamped only when a ticket is first appended."""
    rows = tab.read_all()
    if not rows or not rows[0] or rows[0][0] == "":
        raise RuntimeError("header row missing in sheet")
    if len(rows[0]) < live_col or not rows[0][live_col - 1].strip():
        tab.update_block(1, live_col, [live_hdr])
    if extra and (len(rows[0]) < extra[0] or not rows[0][extra[0] - 1].strip()):
        tab.update_block(1, extra[0], [[extra[1]]])
    pos = {}
    for i, r in enumerate(rows, 1):  # first occurrence wins if a ticket appears twice
        if i > 1 and r and r[0].strip():
            pos.setdefault(nid(r[0]), i)
    byid = {r["TICKET_ID"]: r for r in fresh}

    # 1) refresh live columns of rows already on the sheet (in place)
    ids = [nid(rows[i - 1][0]) if i - 1 < len(rows) else "" for i in range(2, next_row(rows))]
    if ids:
        old = []
        for i in range(2, next_row(rows)):
            row = rows[i - 1] if i - 1 < len(rows) else []
            old.append([(row[live_col - 1 + k] if len(row) > live_col - 1 + k else "") for k in range(len(live_hdr))])
        block = [live_vals(byid[t], o) if t in byid else o for t, o in zip(ids, old)]
        tab.update_block(2, live_col, block)

    # 1b) self-heal: fill any blank fixed cell (e.g. kapture id learned later, or a half-written row)
    for t, i in pos.items():
        if t not in byid:
            continue
        want = [str(v) for v in fixed_vals(byid[t])]
        row = rows[i - 1]
        for j, w in enumerate(want):
            cur = row[j] if len(row) > j else ""
            if not cur.strip() and w.strip():
                tab.update_block(i, j + 1, [[w]])

    # 2) append brand-new tickets at the bottom
    new = [r for r in fresh if r["TICKET_ID"] not in pos]
    if new:
        start = next_row(rows)
        tab.update_block(start, 1, [fixed_vals(r) for r in new])
        tab.update_block(start, live_col, [live_vals(r, [""] * len(live_hdr)) for r in new])
        if extra:
            tab.update_block(start, extra[0], [[extra[2]] for _ in new])
    return len(new)


def main():
    cfg = load_cfg()
    t = datetime.now(IST)
    if os.environ.get("ENFORCE_HOURS") and not (600 <= t.hour * 60 + t.minute <= 1170):
        log("outside 10:00-19:30 IST, skipping")
        return
    if "--check" in sys.argv:  # connectivity test: opens both tabs, writes nothing
        for which in ("caller", "register"):
            log(f"check ok: {which} tab has {len(open_tab(cfg, which).read_all())} rows")
        return
    now = f"{datetime.now(IST):%Y-%m-%d %H:%M}"
    rows = fetch(cfg)
    if not rows:
        log("WARNING: query returned 0 rows - model may have changed, tell Adit. Sheets not touched.")
        sys.exit(2)

    calls = [r for r in rows if r["ARM"] == "CALL"]
    n1 = sync(open_tab(cfg, "caller"), calls,
              lambda r: [r[k] for k in CALLER_FIXED],
              lambda r, old: [r["STATUS_NOW"], r["ISSUE"], now], CALLER_LIVE_COL, CALLER_LIVE_HDR,
              extra=(CALLER_ADDED_COL, "added_at", now))

    def reg_fixed(r):
        return [r["TICKET_ID"], r["KAPTURE_TICKET_ID"], r["CSP_ID"], r["CSP_NAME"], r["ARM"], r["SCORE_BAND"],
                str(r["DELAY_SCORE"]), MODEL, r["CREATED_AT_IST"], r["FLAGGED_AT_IST"], now]

    # register: columns A-K written once, L-N refreshed (status_now, status_updated_at, issue)
    def reg_live(r, old):
        changed = old[0] != r["STATUS_NOW"]
        return [r["STATUS_NOW"], now if changed or not old[1] else old[1], r["ISSUE"]]

    n2 = sync(open_tab(cfg, "register"), rows, reg_fixed, reg_live, REG_LIVE_COL, REG_HDR[REG_LIVE_COL - 1:])
    log(f"ok: {len(rows)} flagged, {len(calls)} CALL; appended {n1} to caller sheet, {n2} to register")


if __name__ == "__main__":
    main()
